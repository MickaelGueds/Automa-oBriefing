import json
from docx import Document

def extract_n25_data(doc_path):
    doc = Document(doc_path)
    commitments_data = {}

    for table in doc.tables:
        
        if "PAUTA DOS COMPROMISSO DO PLANO DE GESTÃO" in table.cell(0, 0).text:
            #
            for row in table.rows[1:]:
                
                cg = row.cells[0].text.strip()

                
                observations = row.cells[1].text.strip()

               
                commitments_data[cg] = {
                    "observations": observations
                }

    return commitments_data
    pass

def save_data_to_json(data, file_name):
    with open(file_name, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=4)

# Exemplo de uso
# doc_path = r'c:\Users\Mickael\Downloads\ATA DA REUNIÃO DE MONITORAMENTO N2.5 (Ciclo 3) - INVESTE.docx'  # Substitua pelo caminho real do documento
# n25_data = extract_n25_data(doc_path)

# save_data_to_json(n25_data, 'n25_data.json')


# for cg, data in n25_data.items():
#     print(f"CG: {cg}")
#     print("Observações:")
#     print(f" - {data['observations']}")
